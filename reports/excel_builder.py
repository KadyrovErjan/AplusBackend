import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, date
from urllib.parse import quote
from django.http import HttpResponse


def parse_date_param(value):
    """Парсит строку даты в формате YYYY-MM-DD или DD-MM-YYYY. Возвращает date или None."""
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


class ExcelBuilder:
    """
    Утилита для генерации Excel-отчетов с единым стилем.

    Использование:
        builder = ExcelBuilder("Заправки", "70AD47", request)
        builder.write_headers(["Дата", "Техника", ...])
        builder.write_row(["01.01.2025", "Хово", ...])
        builder.add_totals_row([5, 6], label_col=2, label="ИТОГО")
        builder.auto_fit_columns()
        builder.freeze_header()
        builder.add_auto_filter()
        return builder.to_response("fuel_report.xlsx")
    """

    def __init__(self, sheet_title, header_color="4F81BD", request=None):
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = sheet_title
        self._request = request
        self._current_row = 1

        self._header_font = Font(bold=True, color="FFFFFF")
        self._header_fill = PatternFill(
            start_color=header_color, end_color=header_color, fill_type="solid"
        )
        self._thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        self._center = Alignment(horizontal="center", vertical="center")
        self._wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
        self._hyperlink_font = Font(color="0563C1", underline="single")
        self._bold_font = Font(bold=True)

    @property
    def current_row(self):
        return self._current_row

    def write_headers(self, headers):
        for col_num, header in enumerate(headers, 1):
            cell = self.ws.cell(row=1, column=col_num, value=header)
            cell.font = self._header_font
            cell.fill = self._header_fill
            cell.alignment = self._center
            cell.border = self._thin_border
        self._current_row = 2

    def write_row(self, values, number_formats=None, center_columns=None, text_columns=None):
        """
        Записывает строку данных.
        number_formats: dict {col_index(1-based): format_str}, например {5: '#,##0.00'}
        center_columns: set of 1-based column indices to center-align
        text_columns: set of 1-based column indices to force text cell type
        """
        row = self._current_row
        for col_num, val in enumerate(values, 1):
            is_text_column = bool(text_columns and col_num in text_columns)
            cell_value = str(val) if is_text_column and val is not None else val
            cell = self.ws.cell(row=row, column=col_num, value=cell_value)
            cell.border = self._thin_border
            if center_columns and col_num in center_columns:
                cell.alignment = self._center
            if is_text_column:
                cell.data_type = "s"
                cell.number_format = "@"
            elif number_formats and col_num in number_formats:
                cell.number_format = number_formats[col_num]
        self._current_row += 1
        return row

    def write_hyperlink_cell(self, row, col, url, label="Ссылка на фото"):
        cell = self.ws.cell(row=row, column=col)
        if url and "http" in str(url):
            cell.hyperlink = url
            cell.value = label
            cell.font = self._hyperlink_font
        else:
            cell.value = url or "Нет фото"
        cell.border = self._thin_border

    def build_absolute_url(self, relative_url):
        if not relative_url:
            return None
        if relative_url.startswith("http"):
            return relative_url
        if self._request:
            return self._request.build_absolute_uri(relative_url)
        return relative_url

    def add_totals_row(self, sum_columns, label_col=1, label="ИТОГО"):
        """
        Добавляет итоговую строку с SUM-формулами.
        sum_columns: list of 1-based column indices для SUM
        """
        row = self._current_row
        data_start = 2
        data_end = row - 1

        if data_end < data_start:
            return

        label_cell = self.ws.cell(row=row, column=label_col, value=label)
        label_cell.font = self._bold_font
        label_cell.border = self._thin_border

        for col in sum_columns:
            col_letter = get_column_letter(col)
            cell = self.ws.cell(
                row=row,
                column=col,
                value=f"=SUM({col_letter}{data_start}:{col_letter}{data_end})",
            )
            cell.font = self._bold_font
            cell.border = self._thin_border
            cell.number_format = "#,##0.00"

        self._current_row += 1

    def auto_fit_columns(self, max_width=50):
        for col in self.ws.columns:
            max_length = 0
            column_letter = col[0].column_letter
            for cell in col:
                try:
                    length = len(str(cell.value))
                    if length > max_length:
                        max_length = length
                except (TypeError, AttributeError):
                    pass
            self.ws.column_dimensions[column_letter].width = min(max_length + 3, max_width)

    def freeze_header(self):
        self.ws.freeze_panes = "A2"

    def add_auto_filter(self):
        if self.ws.dimensions:
            self.ws.auto_filter.ref = self.ws.dimensions

    def to_response(self, filename):
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        encoded_filename = quote(filename)
        response["Content-Disposition"] = (
            f"attachment; filename=report.xlsx; filename*=UTF-8''{encoded_filename}"
        )
        self.wb.save(response)
        return response
