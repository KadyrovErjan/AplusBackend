from datetime import datetime, time, timedelta
from io import BytesIO
from unittest.mock import patch
from urllib.parse import quote

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone

from fuels.models import FuelEvent, FuelReport, FuelReportPhoto, FuelZone
from maintenance.models import IssueReport
from fuels.services.zone_filters import NO_ZONE_ID
from trips.models import Trip
from vehicles.models import DriverShift, Vehicle, VehicleAssignment


class ReportsExportTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.driver = User.objects.create_user(
            email="report_driver@example.com",
            password="test123",
            first_name="Асан",
            surname="Усонов",
            role="driver",
            status="active",
            telegram_id=777001,
        )
        self.vehicle = Vehicle.objects.create(
            veh_id=7539,
            imei="imei-reports-7539",
            name="Хово",
            plate_number="01KG603BO",
            type="Lorry",
        )

    def _load_workbook(self, response):
        return openpyxl.load_workbook(BytesIO(response.content))

    @patch("trips.services.RouteAnalyticsService.get_verified_route_stats")
    def test_trips_export_includes_object_work_type_and_note_columns(self, mock_get_verified_route_stats):
        first_day = datetime(2025, 2, 12).date()
        second_day = datetime(2026, 3, 13).date()
        first_start = timezone.make_aware(datetime.combine(first_day, time(9, 30)))
        first_end = first_start + timedelta(hours=8)
        second_start = timezone.make_aware(datetime.combine(second_day, time(8, 0)))
        second_end = second_start + timedelta(hours=8)
        ignored_start = timezone.make_aware(datetime.combine(datetime(2025, 7, 1).date(), time(7, 45)))
        ignored_end = ignored_start + timedelta(hours=7)

        zone_from = FuelZone.objects.create(name="Объект А", external_id=1952)
        zone_to = FuelZone.objects.create(name="Объект Б", external_id=2000)

        first_shift = DriverShift.objects.create(
            driver=self.driver,
            vehicle=self.vehicle,
            started_at=first_start,
            movement_started_at=first_start,
            ended_at=first_end,
            status="completed",
            trips_count=3,
            distance_km=10.0,
            trips_note="Погрузка с задержкой",
        )
        second_shift = DriverShift.objects.create(
            driver=self.driver,
            vehicle=self.vehicle,
            started_at=second_start,
            movement_started_at=second_start,
            ended_at=second_end,
            status="completed",
            trips_count=4,
            distance_km=12.0,
            trips_note="Второй рейс",
        )
        ignored_shift = DriverShift.objects.create(
            driver=self.driver,
            vehicle=self.vehicle,
            started_at=ignored_start,
            movement_started_at=ignored_start,
            ended_at=ignored_end,
            status="completed",
            trips_count=5,
            distance_km=14.0,
            trips_note="Не нужен в отчете",
        )
        mock_get_verified_route_stats.return_value = {
            first_shift.id: {"count": 3, "distance_km": 0.0, "validated": True, "error": None},
            second_shift.id: {"count": 4, "distance_km": 0.0, "validated": True, "error": None},
            ignored_shift.id: {"count": 0, "distance_km": 0.0, "validated": True, "error": None},
        }

        response = self.client.get(
            "/api/reports/trips/",
            {
                "date_from": first_day.strftime("%d-%m-%Y"),
                "date_to": second_day.strftime("%d-%m-%Y"),
                "zone_from": str(zone_from.external_id),
                "zone_to": str(zone_to.external_id),
            },
        )
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active

        headers = [ws.cell(row=1, column=i).value for i in range(1, 14)]
        self.assertEqual(
            headers,
            [
                "Дата",
                "Техника",
                "Госномер",
                "Водитель",
                "Объект (откуда → куда)",
                "Вид работы",
                "Примечание",
                "Рейсы (водитель)",
                "Рейсы (Incomtek)",
                "Опоздание",
                "Отработано",
                "Норма",
                "Переработка/Недоработка",
            ],
        )

        self.assertEqual(ws.max_row, 4)
        self.assertEqual(ws.cell(row=2, column=1).value, "12.02.2025")
        self.assertEqual(ws.cell(row=3, column=1).value, "13.03.2026")
        self.assertEqual(ws.cell(row=2, column=5).value, "Объект А - Объект Б")
        self.assertEqual(ws.cell(row=3, column=5).value, "Объект А - Объект Б")
        self.assertIsNone(ws.cell(row=2, column=6).value)
        self.assertEqual(ws.cell(row=2, column=7).value, "Погрузка с задержкой")
        self.assertEqual(ws.cell(row=2, column=10).data_type, "s")
        self.assertEqual(ws.cell(row=2, column=11).data_type, "s")
        self.assertEqual(ws.cell(row=2, column=12).data_type, "s")
        self.assertEqual(ws.cell(row=2, column=13).data_type, "s")
        self.assertEqual(ws.cell(row=2, column=10).number_format, "@")
        self.assertEqual(ws.cell(row=2, column=11).number_format, "@")
        self.assertEqual(ws.cell(row=2, column=12).number_format, "@")
        self.assertEqual(ws.cell(row=2, column=13).number_format, "@")

        totals_row = ws.max_row
        self.assertEqual(ws.cell(row=totals_row, column=2).value, "ИТОГО")
        self.assertIsNone(ws.cell(row=totals_row, column=10).value)
        self.assertIsNone(ws.cell(row=totals_row, column=11).value)
        self.assertIsNone(ws.cell(row=totals_row, column=12).value)
        self.assertIsNone(ws.cell(row=totals_row, column=13).value)

    def test_fuels_export_expands_photo_columns(self):
        report = FuelReport.objects.create(
            driver=self.driver,
            vehicle=self.vehicle,
            liters=25.0,
            liters_from_driver=25.0,
            source=FuelReport.Source.MANUAL,
            confirmed=True,
            zone_name="Объект Тест",
        )
        FuelReportPhoto.objects.create(
            report=report,
            photo="fuel_reports/f1.jpg",
            photo_file_id="file_1",
        )
        FuelReportPhoto.objects.create(
            report=report,
            photo="fuel_reports/f2.jpg",
            photo_file_id="file_2",
        )

        today = timezone.localdate()
        response = self.client.get(
            "/api/reports/fuels/",
            {
                "date_from": today.strftime("%d-%m-%Y"),
                "date_to": today.strftime("%d-%m-%Y"),
            },
        )
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active

        expected_headers = [
            "Дата",
            "Техника",
            "Госномер",
            "Водитель",
            "Объект",
            "Вид заправки",
            "ПН-ШНОС",
            "Заправка Incomtek (л)",
            "Заправка водитель (л)",
            "Фото 1",
            "Фото 2",
        ]
        headers = [ws.cell(row=1, column=i).value for i in range(1, len(expected_headers) + 1)]
        self.assertEqual(headers, expected_headers)
        self.assertEqual(ws.cell(row=2, column=5).value, "Объект Тест")
        self.assertEqual(ws.cell(row=2, column=6).value, "дт")
        self.assertIsInstance(ws.cell(row=2, column=1).value, datetime)
        self.assertEqual(ws.cell(row=2, column=1).number_format, "DD.MM.YYYY HH:MM")
        self.assertEqual(ws.cell(row=2, column=10).value, "Фото 1")
        self.assertEqual(ws.cell(row=2, column=11).value, "Фото 2")

    def test_fuels_export_outputs_driver_name_and_username_fallback(self):
        User = get_user_model()
        username_driver = User.objects.create_user(
            email="fuel_username_driver@example.com",
            password="test123",
            username="fuel_driver_login",
            first_name="",
            surname="",
            role="driver",
            status="active",
            telegram_id=777991,
        )

        FuelReport.objects.create(
            driver=self.driver,
            vehicle=self.vehicle,
            liters=25.0,
            liters_from_driver=25.0,
            source=FuelReport.Source.MANUAL,
            confirmed=True,
            zone_name="РђР—РЎ-100",
        )
        FuelReport.objects.create(
            driver=username_driver,
            vehicle=self.vehicle,
            liters=18.0,
            liters_from_driver=18.0,
            source=FuelReport.Source.MANUAL,
            confirmed=True,
            zone_name="РђР—РЎ-500",
        )

        today = timezone.localdate()
        response = self.client.get(
            "/api/reports/fuels/",
            {
                "date_from": today.strftime("%d-%m-%Y"),
                "date_to": today.strftime("%d-%m-%Y"),
            },
        )
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active

        self.assertEqual(ws.cell(row=2, column=4).value, "РђСЃР°РЅ РЈСЃРѕРЅРѕРІ")
        self.assertEqual(ws.cell(row=3, column=4).value, "fuel_driver_login")

    def test_fuels_export_uses_vehicle_driver_without_confirmation(self):
        vehicle_driver = get_user_model().objects.create_user(
            email="vehicle_driver@example.com",
            password="test123",
            first_name="Р‘Р°РєС‹С‚",
            surname="РўСѓСЂСЃСѓРЅРѕРІ",
            role="driver",
            status="active",
            telegram_id=777550,
        )
        self.vehicle.driver = vehicle_driver
        self.vehicle.save(update_fields=["driver"])

        FuelEvent.objects.create(
            vehicle=self.vehicle,
            liters_from_ink=15.0,
            address="РђР—РЎ РЅР° С‚СЂР°СЃСЃРµ",
            zone_name="[]",
            created_at=timezone.make_aware(datetime.combine(timezone.localdate(), time(11, 25))),
        )

        today = timezone.localdate()
        response = self.client.get(
            "/api/reports/fuels/",
            {
                "date_from": today.strftime("%d-%m-%Y"),
                "date_to": today.strftime("%d-%m-%Y"),
            },
        )
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active

        self.assertEqual(ws.cell(row=2, column=4).value, "Р‘Р°РєС‹С‚ РўСѓСЂСЃСѓРЅРѕРІ")

    def test_vehicles_export_uses_row_number_split_name_and_total_mileage_and_assignment_driver(self):
        self.vehicle.mileage = 500
        self.vehicle.total_engine_hours = 100.5
        self.vehicle.last_service_hours = 80.0
        self.vehicle.save(update_fields=["mileage", "total_engine_hours", "last_service_hours"])
        VehicleAssignment.objects.create(vehicle=self.vehicle, driver=self.driver)

        split_vehicle = Vehicle.objects.create(
            veh_id=9102,
            imei="imei-vehicles-9102",
            name="РЎРїСЂРёРЅС‚РµСЂ Р“СЂСѓР·_01KG919AY",
            plate_number="",
            type="Van",
            mileage=13295,
            total_engine_hours=411.2,
            last_service_hours=390.0,
        )
        VehicleAssignment.objects.create(vehicle=split_vehicle, driver=self.driver)

        today = timezone.localdate()
        response = self.client.get(
            "/api/reports/vehicles/",
            {
                "date_from": today.strftime("%d-%m-%Y"),
                "date_to": today.strftime("%d-%m-%Y"),
            },
        )
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active

        headers = [ws.cell(row=1, column=i).value for i in range(1, 8)]
        self.assertEqual(
            headers,
            ["№", "Техника", "Госномер", "Водитель", "Пробег за период (км)", "Моточасы (всего)", "Последнее ТО"],
        )
        self.assertEqual(ws.max_column, 7)

        self.assertEqual(ws.cell(row=2, column=1).value, 1)
        self.assertEqual(ws.cell(row=2, column=2).value, "РҐРѕРІРѕ")
        self.assertEqual(ws.cell(row=2, column=3).value, "01KG603BO")
        self.assertEqual(ws.cell(row=2, column=4).value, "Р С’РЎРѓР В°Р Р… Р Р€РЎРѓР С•Р Р…Р С•Р Р†")
        self.assertEqual(ws.cell(row=2, column=5).value, 500)

        self.assertEqual(ws.cell(row=3, column=1).value, 2)
        self.assertEqual(ws.cell(row=3, column=2).value, "РЎРїСЂРёРЅС‚РµСЂ Р“СЂСѓР·")
        self.assertEqual(ws.cell(row=3, column=3).value, "01KG919AY")
        self.assertEqual(ws.cell(row=3, column=4).value, "РђСЃР°РЅ РЈСЃРѕРЅРѕРІ")
        self.assertEqual(ws.cell(row=3, column=5).value, split_vehicle.mileage)
        self.assertEqual(ws.cell(row=3, column=6).value, split_vehicle.total_engine_hours)
        self.assertEqual(ws.cell(row=3, column=7).value, split_vehicle.last_service_hours)

        totals_row = ws.max_row
        self.assertEqual(ws.cell(row=totals_row, column=2).value, "ИТОГО")

    def test_fuels_export_splits_vehicle_and_plate(self):
        vehicle = Vehicle.objects.create(
            veh_id=9001,
            imei="imei-reports-9001",
            name="Экскаватор 140_01KG381CL",
            plate_number="",
            type="Excavator",
        )
        FuelReport.objects.create(
            driver=self.driver,
            vehicle=vehicle,
            liters=30.0,
            liters_from_driver=30.0,
            source=FuelReport.Source.MANUAL,
            confirmed=True,
            zone_name="АЗС-100",
        )

        today = timezone.localdate()
        response = self.client.get(
            "/api/reports/fuels/",
            {
                "date_from": today.strftime("%d-%m-%Y"),
                "date_to": today.strftime("%d-%m-%Y"),
            },
        )
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active

        self.assertEqual(ws.cell(row=2, column=2).value, "Экскаватор 140")
        self.assertEqual(ws.cell(row=2, column=3).value, "01KG381CL")
        self.assertEqual(ws.cell(row=2, column=5).value, "АЗС-100")

    def test_fuels_export_filters_by_multiple_objects_and_outside(self):
        zone_a = FuelZone.objects.create(name="АЗС-100", external_id=3100)
        zone_b = FuelZone.objects.create(name="АЗС-200", external_id=3200)

        FuelReport.objects.create(
            driver=self.driver,
            vehicle=self.vehicle,
            liters=20.0,
            liters_from_driver=20.0,
            source=FuelReport.Source.MANUAL,
            confirmed=True,
            zone_id=zone_a.id,
            zone_name=zone_a.name,
        )
        FuelReport.objects.create(
            driver=self.driver,
            vehicle=self.vehicle,
            liters=22.0,
            liters_from_driver=22.0,
            source=FuelReport.Source.MANUAL,
            confirmed=True,
            zone_id=zone_b.id,
            zone_name=zone_b.name,
        )
        FuelReport.objects.create(
            driver=self.driver,
            vehicle=self.vehicle,
            liters=18.0,
            liters_from_driver=18.0,
            source=FuelReport.Source.MANUAL,
            confirmed=True,
            zone_id=None,
            zone_name="Вне объекта",
        )

        today = timezone.localdate()
        response = self.client.get(
            "/api/reports/fuels/",
            {
                "date_from": today.strftime("%d-%m-%Y"),
                "date_to": today.strftime("%d-%m-%Y"),
                "zone_ids": [str(zone_a.id), str(NO_ZONE_ID)],
            },
        )
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active

        rows = []
        for row_idx in range(2, ws.max_row):
            rows.append({
                "object": ws.cell(row=row_idx, column=5).value,
            })

        self.assertEqual(len(rows), 2)
        self.assertEqual({row["object"] for row in rows}, {"АЗС-100", "Вне объекта"})

    def test_fuels_export_returns_400_for_invalid_zone_ids(self):
        today = timezone.localdate()
        response = self.client.get(
            "/api/reports/fuels/",
            {
                "date_from": today.strftime("%d-%m-%Y"),
                "date_to": today.strftime("%d-%m-%Y"),
                "zone_ids": "bad-id",
            },
        )
        self.assertEqual(response.status_code, 400)

    def test_fuels_export_object_fallbacks_to_address_when_outside(self):
        self.vehicle.address = "Аква Макс, 7 Апреля улица, Ачекей, Аламудунский район"
        self.vehicle.save(update_fields=["address"])

        FuelReport.objects.create(
            driver=self.driver,
            vehicle=self.vehicle,
            liters=11.0,
            liters_from_driver=11.0,
            source=FuelReport.Source.MANUAL,
            confirmed=True,
            zone_name="Вне объекта",
            address=None,
        )

        today = timezone.localdate()
        response = self.client.get(
            "/api/reports/fuels/",
            {
                "date_from": today.strftime("%d-%m-%Y"),
                "date_to": today.strftime("%d-%m-%Y"),
            },
        )
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active
        self.assertEqual(
            ws.cell(row=2, column=5).value,
            "Аква Макс, 7 Апреля улица, Ачекей, Аламудунский район",
        )

    def test_fuels_export_uses_fueling_address_when_zone_name_is_empty_list(self):
        self.vehicle.address = "Старый адрес техники"
        self.vehicle.save(update_fields=["address"])

        FuelReport.objects.create(
            driver=self.driver,
            vehicle=self.vehicle,
            liters=14.0,
            liters_from_driver=14.0,
            source=FuelReport.Source.MANUAL,
            confirmed=True,
            zone_name="[]",
            address="Проспект Чуй, 1",
        )

        today = timezone.localdate()
        response = self.client.get(
            "/api/reports/fuels/",
            {
                "date_from": today.strftime("%d-%m-%Y"),
                "date_to": today.strftime("%d-%m-%Y"),
            },
        )
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active
        self.assertEqual(ws.cell(row=2, column=5).value, "Проспект Чуй, 1")

    def test_fuels_export_merges_incomtek_and_driver_rows_by_vehicle_number(self):
        event_dt = timezone.make_aware(datetime.combine(timezone.localdate(), time(12, 41)))
        FuelEvent.objects.create(
            vehicle=self.vehicle,
            liters_from_ink=32.5,
            address="Тестовый адрес заправки",
            zone_name="[]",
            created_at=event_dt,
        )

        report = FuelReport.objects.create(
            driver=self.driver,
            vehicle=self.vehicle,
            liters=30.0,
            liters_from_driver=30.0,
            source=FuelReport.Source.MANUAL,
            confirmed=True,
            zone_name="Вне объекта",
            address="Тестовый адрес заправки",
        )
        FuelReport.objects.filter(pk=report.pk).update(created_at=event_dt)
        report.refresh_from_db()

        FuelReportPhoto.objects.create(
            report=report,
            photo="fuel_reports/merged.jpg",
            photo_file_id="merged_1",
        )

        today = timezone.localdate()
        response = self.client.get(
            "/api/reports/fuels/",
            {
                "date_from": today.strftime("%d-%m-%Y"),
                "date_to": today.strftime("%d-%m-%Y"),
            },
        )
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active

        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.cell(row=2, column=1).value, event_dt.replace(tzinfo=None))
        self.assertEqual(ws.cell(row=2, column=1).number_format, "DD.MM.YYYY HH:MM")
        self.assertEqual(ws.cell(row=2, column=3).value, "01KG603BO")
        self.assertEqual(ws.cell(row=2, column=5).value, "Тестовый адрес заправки")
        self.assertEqual(ws.cell(row=2, column=8).value, 32.5)
        self.assertEqual(ws.cell(row=2, column=9).value, 30.0)
        self.assertEqual(ws.cell(row=2, column=10).value, "Фото 1")

    def test_fuels_export_keeps_multiple_fuelings_for_same_vehicle(self):
        first_dt = timezone.make_aware(datetime.combine(timezone.localdate() - timedelta(days=20), time(14, 10, 3)))
        second_dt = timezone.make_aware(datetime.combine(timezone.localdate(), time(14, 9, 17)))

        FuelEvent.objects.create(
            vehicle=self.vehicle,
            liters_from_ink=20.0,
            address="Первая заправка",
            zone_name="[]",
            created_at=first_dt,
        )
        FuelEvent.objects.create(
            vehicle=self.vehicle,
            liters_from_ink=25.0,
            address="Вторая заправка",
            zone_name="[]",
            created_at=second_dt,
        )

        response = self.client.get(
            "/api/reports/fuels/",
            {
                "date_from": (timezone.localdate() - timedelta(days=30)).strftime("%d-%m-%Y"),
                "date_to": timezone.localdate().strftime("%d-%m-%Y"),
            },
        )
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active

        self.assertEqual(ws.max_row, 4)
        self.assertEqual(ws.cell(row=2, column=1).value, second_dt.replace(tzinfo=None))
        self.assertEqual(ws.cell(row=3, column=1).value, first_dt.replace(tzinfo=None))
        self.assertEqual(ws.cell(row=2, column=5).value, "Вторая заправка")
        self.assertEqual(ws.cell(row=3, column=5).value, "Первая заправка")

    def test_trips_export_uses_latest_historical_assignment_when_shift_vehicle_missing(self):
        first_vehicle = Vehicle.objects.create(
            veh_id=9301,
            imei="imei-trips-9301",
            name="РЎС‚Р°СЂР°СЏ С‚РµС…РЅРёРєР°",
            plate_number="01KG001AA",
            type="Lorry",
        )
        second_vehicle = Vehicle.objects.create(
            veh_id=9302,
            imei="imei-trips-9302",
            name="РќРѕРІР°СЏ С‚РµС…РЅРёРєР°",
            plate_number="01KG002BB",
            type="Lorry",
        )
        shift_day = timezone.localdate() - timedelta(days=1)
        shift_started_at = timezone.make_aware(datetime.combine(shift_day, time(9, 0)))
        shift_ended_at = shift_started_at + timedelta(hours=8)

        first_assignment = VehicleAssignment.objects.create(vehicle=first_vehicle, driver=self.driver)
        second_assignment = VehicleAssignment.objects.create(vehicle=second_vehicle, driver=self.driver)
        VehicleAssignment.objects.filter(pk=first_assignment.pk).update(
            assigned_at=shift_started_at - timedelta(days=2),
            unassigned_at=shift_started_at + timedelta(days=1),
        )
        VehicleAssignment.objects.filter(pk=second_assignment.pk).update(
            assigned_at=shift_started_at - timedelta(hours=2),
            unassigned_at=shift_started_at + timedelta(days=2),
        )

        DriverShift.objects.create(
            driver=self.driver,
            vehicle=None,
            started_at=shift_started_at,
            movement_started_at=shift_started_at,
            ended_at=shift_ended_at,
            status="completed",
            trips_count=2,
            distance_km=8.5,
            trips_note="РЎРјРµРЅР° Р±РµР· СЏРІРЅРѕР№ СЃРІСЏР·Рё СЃ С‚РµС…РЅРёРєРѕР№",
        )

        response = self.client.get(
            "/api/reports/trips/",
            {
                "date_from": shift_day.strftime("%d-%m-%Y"),
                "date_to": shift_day.strftime("%d-%m-%Y"),
                "plate_number": "01KG002BB",
            },
        )
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active

        self.assertEqual(ws.cell(row=2, column=2).value, "РќРѕРІР°СЏ С‚РµС…РЅРёРєР°")
        self.assertEqual(ws.cell(row=2, column=3).value, "01KG002BB")
        self.assertEqual(ws.cell(row=2, column=8).value, 2)

    @patch("trips.services.RouteAnalyticsService.get_verified_route_stats")
    def test_trips_export_defaults_date_to_today_and_uses_local_trip_matches(self, mock_get_verified_route_stats):
        today = timezone.localdate()
        tomorrow = today + timedelta(days=1)
        current_start = timezone.make_aware(datetime.combine(today, time(9, 0)))
        current_end = current_start + timedelta(hours=8)
        future_start = timezone.make_aware(datetime.combine(tomorrow, time(9, 0)))
        future_end = future_start + timedelta(hours=8)

        zone_from = FuelZone.objects.create(name="Аристократ", external_id=1952)
        zone_to = FuelZone.objects.create(name="Полигон", external_id=2000)

        current_shift = DriverShift.objects.create(
            driver=self.driver,
            vehicle=self.vehicle,
            started_at=current_start,
            movement_started_at=current_start,
            ended_at=current_end,
            status="completed",
            trips_count=2,
            distance_km=10.0,
            trips_note="Есть локальный Trip",
        )
        DriverShift.objects.create(
            driver=self.driver,
            vehicle=self.vehicle,
            started_at=future_start,
            movement_started_at=future_start,
            ended_at=future_end,
            status="completed",
            trips_count=5,
            distance_km=20.0,
            trips_note="Будущая смена не должна войти",
        )
        Trip.objects.create(
            vehicle=self.vehicle,
            start_zone_name=zone_from.name,
            end_zone_name=zone_to.name,
            start_at=current_start + timedelta(minutes=15),
            arrival_at=current_start + timedelta(minutes=35),
            end_at=current_start + timedelta(minutes=45),
            trip_duration_seconds=1200,
            dwell_duration_seconds=600,
            distance_km=2.5,
        )

        response = self.client.get(
            "/api/reports/trips/",
            {
                "date_from": today.strftime("%d-%m-%Y"),
                "zone_from": str(zone_from.external_id),
                "zone_to": str(zone_to.external_id),
            },
        )
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active

        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.cell(row=2, column=1).value, today.strftime("%d.%m.%Y"))
        self.assertEqual(ws.cell(row=2, column=7).value, current_shift.trips_note)
        self.assertEqual(ws.cell(row=2, column=9).value, 1)
        mock_get_verified_route_stats.assert_not_called()


    def test_fuels_export_sorts_dates_in_ascending_order(self):
        first_dt = timezone.make_aware(datetime.combine(timezone.localdate() - timedelta(days=2), time(8, 0)))
        second_dt = timezone.make_aware(datetime.combine(timezone.localdate() - timedelta(days=1), time(9, 30)))

        FuelEvent.objects.create(
            vehicle=self.vehicle,
            liters_from_ink=10.0,
            address="Ранняя заправка",
            zone_name="[]",
            created_at=first_dt,
        )
        FuelEvent.objects.create(
            vehicle=self.vehicle,
            liters_from_ink=12.0,
            address="Поздняя заправка",
            zone_name="[]",
            created_at=second_dt,
        )

        response = self.client.get(
            "/api/reports/fuels/",
            {
                "date_from": (timezone.localdate() - timedelta(days=3)).strftime("%d-%m-%Y"),
                "date_to": timezone.localdate().strftime("%d-%m-%Y"),
            },
        )
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active

        self.assertEqual(ws.cell(row=2, column=1).value, first_dt.replace(tzinfo=None))
        self.assertEqual(ws.cell(row=3, column=1).value, second_dt.replace(tzinfo=None))

    def test_fuels_export_uses_russian_date_range_filename(self):
        today = timezone.localdate()
        date_from = today - timedelta(days=1)

        response = self.client.get(
            "/api/reports/fuels/",
            {
                "date_from": date_from.strftime("%d-%m-%Y"),
                "date_to": today.strftime("%d-%m-%Y"),
            },
        )
        self.assertEqual(response.status_code, 200)
        expected_filename = (
            "attachment; filename=report.xlsx; filename*=UTF-8''"
            + quote(
                f"отчет_о_заправке_от_{date_from.strftime('%d-%m-%Y')}"
                f"_до_{today.strftime('%d-%m-%Y')}.xlsx"
            )
        )
        self.assertEqual(response["Content-Disposition"], expected_filename)

    def test_repairs_export_outputs_driver_name(self):
        vehicle = Vehicle.objects.create(
            veh_id=9103,
            imei="imei-repairs-driver-9103",
            name="РҐРѕРІРѕ_01KG777AA",
            plate_number="",
            type="Lorry",
        )
        IssueReport.objects.create(
            driver=self.driver,
            vehicle=vehicle,
            description="РџСЂРѕРІРµСЂРєР° РІРѕРґРёС‚РµР»СЏ",
            status="pending",
        )

        response = self.client.get("/api/reports/repairs/")
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active

        self.assertEqual(ws.cell(row=2, column=5).value, "РђСЃР°РЅ РЈСЃРѕРЅРѕРІ")

    def test_repairs_export_defaults_to_today_and_formats_columns(self):
        vehicle = Vehicle.objects.create(
            veh_id=9101,
            imei="imei-repairs-9101",
            name="Спринтер Груз_01KG919AY",
            plate_number="",
            type="Van",
        )

        today_issue = IssueReport.objects.create(
            driver=self.driver,
            vehicle=vehicle,
            description="Нужно заменить колесо",
            status="pending",
        )
        old_issue = IssueReport.objects.create(
            driver=self.driver,
            vehicle=vehicle,
            description="Старый ремонт",
            status="resolved",
        )

        today_dt = timezone.make_aware(datetime.combine(timezone.localdate(), time(10, 0)))
        old_dt = timezone.make_aware(datetime.combine(timezone.localdate() - timedelta(days=2), time(9, 0)))
        IssueReport.objects.filter(pk=today_issue.pk).update(created_at=today_dt)
        IssueReport.objects.filter(pk=old_issue.pk).update(created_at=old_dt)

        response = self.client.get("/api/reports/repairs/")
        self.assertEqual(response.status_code, 200)

        wb = self._load_workbook(response)
        ws = wb.active

        headers = [ws.cell(row=1, column=i).value for i in range(1, 9)]
        self.assertEqual(
            headers,
            ["№", "Дата/Время", "Техника", "Госномер", "Водитель", "Описание проблемы", "Цена", "Фото 1"],
        )
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws.cell(row=2, column=1).value, 1)
        self.assertEqual(ws.cell(row=2, column=3).value, "Спринтер Груз")
        self.assertEqual(ws.cell(row=2, column=4).value, "01KG919AY")
        self.assertEqual(ws.cell(row=2, column=6).value, "Нужно заменить колесо")
        self.assertEqual(ws.cell(row=2, column=7).value, 0)

        expected_filename = (
            "attachment; filename=report.xlsx; filename*=UTF-8''"
            + quote(f"отчет_о_ремонтах_от_{timezone.localdate().strftime('%d-%m-%Y')}"
                    f"_до_{timezone.localdate().strftime('%d-%m-%Y')}.xlsx")
        )
        self.assertEqual(response["Content-Disposition"], expected_filename)
